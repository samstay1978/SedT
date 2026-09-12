#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Office文档敏感词加解密工具 v2.4
支持格式: Word(.docx), Excel(.xlsx), PowerPoint(.pptx)
v2.4 更新内容:
1. 新增"强加密"模式（单文件加密、批量加密均可勾选）：
   - 替换串与敏感词不再等长：敏感词长度为 L 时，替换串长度在 [L, 2L] 间随机
   - 每个敏感词生成 5~10 个不同的英文字符串，文档中每次出现随机选用其一
   - 两个特征同时生效，显著提高从密文反推敏感词的难度
2. 解密完全自动适配：标准/强加密词表均可解密，无需手动选择
v2.3 更新内容:
1. 词表加密升级为 AES-256-GCM（此前为 AES-128-CBC，说明书宣称 AES-256 不实），
   同时兼容解密 v2.2 生成的 Fernet 格式词表
2. 词表与文档绑定：加密时写入加密文档的 SHA-256，解密时校验，防止词表混用
3. Word 覆盖范围扩展：超链接文本、页眉/页脚、文本框内容均可替换
4. Excel 数字/日期单元格类型保护：命中敏感词时保持数值/日期类型，无法保持则跳过并提示
5. 加解密改为后台线程执行，界面不卡顿，操作完成有明确反馈
6. 目标文件已存在时弹出覆盖确认，避免误覆盖
7. 新增批量处理：可对一个目录中的文档批量加密 / 批量解密
"""

import os
import sys
import json
import base64
import hashlib
import secrets
import string
import queue
import threading
import subprocess
import webbrowser
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, scrolledtext
from pathlib import Path
from abc import ABC, abstractmethod

# ==================== 内嵌说明书内容 ====================

USER_MANUAL_CONTENT = """══════════════════════════════════════════════════════════════════
                    Office敏感词加解密工具 — 使用说明书
                              版本: v2.4
══════════════════════════════════════════════════════════════════

【一、软件简介】

本工具支持对 Microsoft Office 三件套（Word、Excel、PowerPoint）
文档中的敏感词汇进行英文字符替换加密，并采用 AES-256-GCM
算法对词表进行加密保护。加密后的文档外观保持正常排版，敏感词被
替换为随机英文字符串，只有通过配套词表和密码才能还原原文。
标准模式使用与敏感词等长的随机字符串；勾选"强加密"后，替换串
长度在 L~2L 间随机且每个敏感词配备 5~10 个不同字符串逐处随机
替换，隐蔽性更强。

支持格式:
  • Word 文档      (.docx)
  • Excel 工作簿   (.xlsx)
  • PowerPoint 演示文稿 (.pptx)

v2.4 新增:
  • 强加密模式：替换串不再等长（长度 L~2L 随机），每个敏感词
    配备 5~10 个不同英文字符串，文档中每次出现随机换用，单文件
    加密与批量加密均可勾选
  • 解密自动适配标准/强加密词表，无需手动选择

v2.3 新增:
  • 批量处理：可对一个目录中的文档批量加密 / 批量解密
  • 词表-文档绑定：词表与加密文档一一对应，混用会提示不匹配
  • Word 超链接、页眉页脚、文本框内容同样支持替换
  • Excel 数字/日期单元格自动保护类型


【二、系统要求】

  • 操作系统: Windows 7 / 8 / 10 / 11
  • 运行环境: 无需安装 Python，独立运行
  • 依赖组件: 程序已自带所有运行库


【三、界面说明】

左侧: 敏感词表管理（两个页签通用）
  • 添加/删除/清空敏感词
  • 加载/保存词表文件 (.json)
  • 生成示例词表

右侧页签一「单文件处理」:
  • 左栏-文档加密: 选择 Office 文档、设置解密密码、
    可勾选"强加密"提升替换复杂度，一键加密，生成
    加密文档 + 加密词表
  • 右栏-文档解密: 选择加密后的文档、配套 .enc 加密词表、
    输入密码，一键还原（自动识别标准/强加密）

右侧页签二「批量处理」:
  • 批量加密: 选择源目录与输出目录，加密整个目录中的文档
  • 批量解密: 选择加密文档目录，自动配对词表并逐一还原
  • 批量加密/解密共用同一密码框

窗口缩放:
  • 窗口可自由缩放，界面为固定布局（无滚动条）
  • 最小窗口尺寸已保证所有控件完整可见，不会被裁切


【四、操作步骤】

▶ 单个加密流程:
  1. 在左侧"敏感词表管理"中添加需要保护的敏感词汇
     （如: 机密、薪酬、密码、身份证号 等）
  2. 点击"加载示例词表"可快速体验功能
  3. 在"文档加密"面板中，点击"浏览..."选择要加密的 Office 文档
  4. 在"密码"框中设置解密密码（请务必牢记）
  5. 如需更强隐蔽性，勾选密码框旁的"强加密"：替换串长度在
     L~2L 间随机，每个敏感词配备 5~10 个不同字符串，文档中
     同一敏感词的每次出现都会被替换成不同的英文字符串
  6. 点击"一键加密"，选择保存位置
  7. 程序将生成两个文件:
       xxx_encrypted.docx  — 加密后的文档
       xxx_vocab.enc       — 加密后的词表（必须妥善保管）

▶ 单个解密流程:
  1. 在"文档解密"面板中，选择加密后的文档（_encrypted 后缀）
  2. 选择配套的 .enc 加密词表文件
  3. 输入加密时设置的密码
  4. 点击"一键解密"，选择保存位置
  5. 程序将恢复原始文档和词表

▶ 批量加密流程:
  1. 在"批量处理"面板中，选择源目录（存放待加密文档的文件夹）
  2. 选择输出目录（留空则输出到源目录）
  3. 输入密码；勾选"强加密（批量加密生效）"则整批文档采用
     变长、多变体随机替换（与单文件强加密规则相同）
  4. 点击"批量加密"
  5. 目录内所有 .docx/.xlsx/.pptx 文件将被逐一加密。
     默认生成一个共享词表 _batch_vocab.enc；
     勾选"每文件独立词表"则每个文件生成配套 _vocab.enc；
     强加密可与上述两种词表模式任意组合

▶ 批量解密流程:
  1. 选择加密文档所在目录
  2. 词表目录默认与加密文档同目录
  3. 输入密码，点击"批量解密"
  4. 程序自动寻找共享词表或按文件名配对词表并逐一还原


【五、重要提示】

  ⚠ 密码是唯一的解密凭证，密码丢失将无法恢复原文！
  ⚠ 加密词表 (.enc) 必须与加密文档配套使用，不可混用；
    v2.3 起程序会校验文档指纹，混用将提示"文档与词表不匹配"。
  ⚠ 建议将加密文档和加密词表分开保管，提高安全性。
  ⚠ 加密后的文档中，敏感词被替换为英文字符串，文档
    排版和格式基本保持不变，可直接用于日常流转。标准模式
    为等长替换；强加密替换串长度为 L~2L，局部行宽可能略有
    变化，但不影响文档打开与正常阅读。
  ⚠ 本工具采用零宽空格标记技术，确保解密时不会误替换文档
    中原本存在的正常英文内容。
  ⚠ v2.3 起词表使用 AES-256-GCM 认证加密；v2.2 生成的词表
    仍可正常解密（Fernet/AES-128 格式兼容）。


【六、桌面图标与右键菜单（打包版）】

  打包为 exe / MSIX 安装后，首次启动程序会自动完成：

  1. 桌面图标
     - 在桌面创建「Office敏感词加解密工具」快捷方式；
     - MSIX 安装后 exe 位于系统受限目录（WindowsApps），普通用户
       无法直接执行该路径；快捷方式因此经 explorer.exe 打开
       shell:AppsFolder\<包系列名>!<应用ID>（系统应用模型激活）
       启动，双击即可正常运行，MSIX 升级后始终有效；
     - MSIX 每次启动都会校验并重建快捷方式，无需手动处理。

  2. 文件右键菜单
     - 对 .docx / .xlsx / .pptx 文件右键，出现两个菜单项：
         「用Office敏感词工具加密」
         「用Office敏感词工具解密」
     - 点击后程序自动打开并选中该文件，输入密码即可操作；
     - 注册写入当前用户注册表 (HKCU)，无需管理员权限；
     - 程序每次启动会校验路径，MSIX 更新后自动重注册。

  3. 命令行用法（高级）
     - OfficeSensitiveEncryptor.exe --encrypt 文件路径
     - OfficeSensitiveEncryptor.exe --decrypt 文件路径
     - 直接拖拽文件到 exe 上启动，等同加密模式

  注：以上功能仅在打包（exe/MSIX）后生效；源码方式运行时
      不会修改桌面与注册表。


【七、常见问题】

Q1: 加密后的文档可以正常打开吗？
A: 可以。加密后的文档外观与原文档一致，只是敏感词变成了
   英文字符串，不影响正常浏览和打印。

Q2: 密码忘记了怎么办？
A: 密码丢失后无法恢复原文，请务必妥善保管。建议将密码
   记录在安全的地方。

Q3: 支持哪些文件格式？
A: 支持 .docx (Word)、.xlsx (Excel)、.pptx (PowerPoint)。
   不支持 .doc、.xls、.ppt 等旧版二进制格式。

Q4: 加密词表可以通用吗？
A: 不可以。每次加密都会生成独立的词表，且 v2.3 起词表与
   文档指纹绑定，不同文档的词表互不通用。解密时必须使用
   与文档配套的 .enc 词表文件。

Q5: Excel 富文本单元格格式会丢失吗？
A: 如果敏感词完全落在单个格式块内，格式完全保持；如果
   敏感词跨越多个格式块，可能合并为纯文本（内容绝对正确）。

Q6: Excel 中的数字会被加密吗？
A: 数字/日期单元格不会参与替换，可避免类型被破坏。若敏感词
   出现在数字单元格中，程序会跳过该单元格并在日志中提示。

Q7: 批量加密支持子目录吗？
A: v2.3 仅处理所选目录的直接文件（不递归子目录）。

Q8: "强加密"和标准模式有什么区别？该怎么选？
A: 标准模式：每个敏感词只对应 1 个等长英文字符串，文中
   所有相同敏感词都被替换成同一个串，替换前后长度不变。
   强加密：敏感词长度为 L 时，替换串长度在 L~2L 之间随机；
   每个敏感词一次性生成 5~10 个不同字符串，文中每次出现
   随机换用其一。这样无法通过"密文出现频率/长度"反推敏感
   词，抗分析能力更强。强加密仅影响加密环节，解密时程序
   自动识别词表类型，操作步骤完全一样；两种模式生成的词表
   互不通用（各自与文档指纹绑定）。

Q9: 强加密后文档排版变化比标准模式明显，正常吗？
A: 正常。强加密的替换串可能比原敏感词长（最长 2 倍），局部
   文字行宽会有细微变化；字体、字号、加粗、颜色等字符格式
   依旧保留，不影响打开、浏览和打印。


【八、技术支持】

  作者: Sam Li
  邮箱: samstay@sina.com
  如有问题或建议，欢迎通过邮件联系。

  请访问项目网站获取最新版本和更多功能：https://samstay.dpdns.org/

══════════════════════════════════════════════════════════════════
"""

COPYRIGHT_CONTENT = """══════════════════════════════════════════════════════════════════
                    Office敏感词加解密工具 — 版权声明与授权协议
                              版本: v2.4
                         版权所有 © Sam Li
                         联系邮箱: samstay@sina.com
══════════════════════════════════════════════════════════════════

【一、版权声明】

  本软件"Office敏感词加解密工具"（以下简称"本软件"）由 Sam Li
  独立开发完成，受《中华人民共和国著作权法》、《计算机软件
  保护条例》及其他相关法律法规保护。

  版权所有: Sam Li
  联系邮箱: samstay@sina.com
  未经授权，任何单位或个人不得以任何形式复制、修改、传播、
  出售或用于商业目的。


【二、授权许可】

▶ 个人非商业用途 — 免费授权
  • 个人用户可在非商业场景下免费使用本软件的全部功能。
  • 包括但不限于: 个人学习、研究、家庭文档管理、个人项目
    等非盈利性活动。
  • 个人用户可自由复制和分发本软件，但须保持软件完整性，
    不得移除或修改本版权声明。

▶ 企业/单位商业用途 — 必须购买商用授权
  • 任何企业、事业单位、组织机构将本软件用于商业生产、
    内部业务流程、对外服务、产品集成、数据安全处理等
    直接或间接产生经济收益的场景，必须事先获得书面商用授权。
  • 包括但不限于以下情形:
    - 企业内部文档加密管理系统的组成部分
    - 对外提供文档安全处理服务
    - 集成到商业产品或解决方案中
    - 政府、金融、医疗、教育等机构的业务系统使用
    - 任何形式的 SaaS、PaaS 或托管服务中使用


【三、禁止行为】

  未经版权所有者 Sam Li 的书面许可，任何单位或个人不得:
  1. 对本软件进行反向工程、反编译、反汇编或试图以其他方式
     发现软件的源代码；
  2. 出租、出借、出售、分发、再许可本软件或其任何部分；
  3. 移除或修改软件中的版权声明、授权协议、作者信息；
  4. 将本软件用于任何违反法律法规或侵犯第三方权益的活动；
  5. 以本软件为基础开发竞争性产品或服务。


【四、免责声明】

  1. 本软件按"现状"提供，作者不对软件的适用性、可靠性、
     准确性作出任何明示或暗示的担保。
  2. 用户应自行承担使用本软件的风险，作者不对因使用或无法
     使用本软件而导致的任何直接、间接、附带、特殊或后果性
     损害承担责任，包括但不限于数据丢失、业务中断、密码
     遗忘导致的文档无法恢复等情形。
  3. 用户有责任妥善保管加密密码和加密词表，因密码丢失或
     词表损坏导致的文档无法解密，作者不提供恢复服务且不
     承担任何责任。
  4. 本软件采用的加密算法（AES-256-GCM）在现有技术条件下具有
     较高的安全性，但不保证能够抵御未来可能出现的密码学
     攻击手段。


【五、授权申请】

  如需购买企业商用授权，请通过以下方式联系:

    联系人: Sam Li
    电子邮箱: samstay@sina.com

  请在邮件中注明:
    • 企业/单位全称
    • 统一社会信用代码（如有）
    • 预计使用场景和范围
    • 联系人及联系电话

  收到申请后，我们将在 3 个工作日内回复授权方案及报价。
  未经授权擅自用于商业用途的，版权所有者保留追究法律
  责任的权利。


【六、协议变更】

  版权所有者保留随时修改本授权协议的权利，修改后的协议
  将在软件更新时一并发布。继续使用本软件即视为接受修改
  后的协议条款。


【七、法律适用与争议解决】

  本授权协议适用中华人民共和国法律。因本协议引起的或与本
  协议有关的任何争议，双方应友好协商解决；协商不成的，
  任何一方均可向版权所有者所在地有管辖权的人民法院提起
  诉讼。

══════════════════════════════════════════════════════════════════
                         感谢您使用本软件！
══════════════════════════════════════════════════════════════════
"""


# ==================== 辅助函数：获取程序目录（兼容打包） ====================

def get_app_dir() -> str:
    """获取程序所在目录，兼容 PyInstaller/Nuitka 打包后的 exe 环境"""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def get_word_dir() -> str:
    """词表目录：程序目录下的 word 子目录（内置领域词表 + 示例文档）。
    不存在则自动创建；加载/保存词表对话框首次打开默认定位到该目录。"""
    d = os.path.join(get_app_dir(), "word")
    try:
        os.makedirs(d, exist_ok=True)
    except Exception:
        d = get_app_dir()
    return d


def ensure_manual_files():
    """程序启动时，将内嵌的说明书内容释放到程序同级目录"""
    app_dir = get_app_dir()
    manual_path = os.path.join(app_dir, "README.txt")
    copyright_path = os.path.join(app_dir, "LICENSE.txt")

    if not os.path.exists(manual_path):
        try:
            # utf-8-sig 带 BOM，Windows 记事本各版本均可正常显示
            with open(manual_path, "w", encoding="utf-8-sig") as f:
                f.write(USER_MANUAL_CONTENT)
        except Exception:
            pass

    if not os.path.exists(copyright_path):
        try:
            with open(copyright_path, "w", encoding="utf-8-sig") as f:
                f.write(COPYRIGHT_CONTENT)
        except Exception:
            pass


def open_text_file(filename: str):
    """使用系统默认程序打开 txt 文件"""
    app_dir = get_app_dir()
    filepath = os.path.join(app_dir, filename)

    if not os.path.exists(filepath):
        messagebox.showerror("错误", f"未找到文件: {filename}")
        return

    try:
        os.startfile(filepath)
    except Exception as e:
        messagebox.showerror("错误", f"无法打开文件: {e}")


# ==================== 桌面图标 / 右键菜单 / 命令行参数（打包环境） ====================

def parse_cli_args():
    """解析命令行参数，支持:
    - 程序 --encrypt <文件>
    - 程序 --decrypt <文件>
    - 程序 <文件>            （自动判断为加密）
    返回 (mode, path)，mode 为 'encrypt'/'decrypt'/None"""
    args = sys.argv[1:]
    mode, path = None, None
    i = 0
    while i < len(args):
        a = args[i]
        if a in ("--encrypt", "--decrypt"):
            mode = a[2:]
            if i + 1 < len(args):
                path = args[i + 1]
                i += 1
        elif a.startswith("-"):
            pass  # 其它开关忽略
        elif path is None and os.path.isfile(a):
            path = a
        i += 1
    if path is None:
        return None, None
    return mode or "encrypt", path


def _run_powershell(script: str):
    """静默运行 PowerShell（无窗口），返回 stdout；失败返回空串"""
    try:
        flags = 0x08000000  # CREATE_NO_WINDOW
        r = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
            capture_output=True, text=True, timeout=30, creationflags=flags,
        )
        return (r.stdout or "").strip()
    except Exception:
        return ""


def _get_package_family_name() -> str:
    """检测当前进程是否具有 MSIX 包身份：有则返回包系列名(PFN)，无则返回空串。
    用于区分 MSIX 安装环境与普通 exe 打包环境。"""
    if not sys.platform.startswith("win"):
        return ""
    try:
        import ctypes
        k32 = ctypes.windll.kernel32
        n = ctypes.c_uint32(0)
        # 传空缓冲区：预期返回 122(ERROR_INSUFFICIENT_BUFFER) 并写出所需字符数；
        # 无包身份时返回 APPMODEL_ERROR_NO_PACKAGE(15700) 等其它值
        if k32.GetPackageFamilyName(k32.GetCurrentProcess(), ctypes.byref(n), None) != 122:
            return ""
        buf = ctypes.create_unicode_buffer(n.value)
        if k32.GetPackageFamilyName(k32.GetCurrentProcess(), ctypes.byref(n), buf) == 0:
            return buf.value
    except Exception:
        pass
    return ""


def ensure_desktop_shortcut():
    """MSIX/打包后首次启动：创建（或修正）桌面快捷方式。
    MSIX 下 exe 本体位于 WindowsApps 受限目录，普通用户进程无法直接
    执行该路径（直接指向 exe 的快捷方式双击报"无法访问"），因此快捷
    方式经 explorer.exe 打开 shell:AppsFolder\\<包系列名>!<应用ID>，
    走系统应用模型激活，任何 Windows 10+ 环境均可用；若检测到每用户
    应用执行别名存在则优先使用别名路径。普通 exe 打包环境仍直接指向
    exe 本体。MSIX 每次启动都重建快捷方式，升级后自动跟随。"""
    if not getattr(sys, "frozen", False) or not sys.platform.startswith("win"):
        return
    try:
        exe = sys.executable
        pfn = _get_package_family_name()
        if pfn:
            alias = os.path.join(
                os.environ.get("LOCALAPPDATA", ""),
                "Microsoft", "WindowsApps", "OfficeSensitiveEncryptor.exe")
            if os.path.exists(alias):
                target, args = alias, ""
                workdir = os.path.dirname(alias)
                icon = alias
            else:
                target = os.path.join(os.environ.get("WINDIR", r"C:\Windows"),
                                      "explorer.exe")
                args = f"shell:AppsFolder\\{pfn}!OfficeSensitiveEncryptor"
                workdir = os.path.dirname(exe)
                icon = exe
        else:
            target, args = exe, ""
            workdir = os.path.dirname(exe)
            icon = exe
        target_q = target.replace("'", "''")
        args_q = args.replace("'", "''")
        workdir_q = workdir.replace("'", "''")
        icon_q = icon.replace("'", "''")
        script = (
            "$d=[Environment]::GetFolderPath('Desktop');"
            "if($d){"
            "$s=(New-Object -ComObject WScript.Shell).CreateShortcut("
            f"[System.IO.Path]::Combine($d,'Office敏感词加解密工具.lnk'));"
            f"$s.TargetPath='{target_q}';"
            f"$s.Arguments='{args_q}';"
            f"$s.WorkingDirectory='{workdir_q}';"
            f"$s.IconLocation='{icon_q},0';"
            "$s.Save()}"
        )
        _run_powershell(script)
    except Exception:
        pass


def ensure_context_menu():
    """MSIX/打包后首次启动：注册（或修正）文件右键菜单。
    对 .docx/.xlsx/.pptx 注册「加密」「解密」两项（HKCU，免管理员）。
    MSIX 更新后 exe 路径变化，每次都校验路径并重注册。"""
    if not getattr(sys, "frozen", False) or not sys.platform.startswith("win"):
        return
    try:
        import winreg
    except Exception:
        return
    exe = sys.executable
    exe_q = exe.replace('"', '""')
    quoted = f'"{exe_q}"'

    def set_menu(ext, action, label):
        base = rf"Software\Classes\SystemFileAssociations\{ext}\shell\OfficeSensitive{action.capitalize()}"
        try:
            winreg.DeleteKey(winreg.HKEY_CURRENT_USER, base + r"\command")  # 旧值先删
        except OSError:
            pass
        try:
            winreg.DeleteKey(winreg.HKEY_CURRENT_USER, base)
        except OSError:
            pass
        try:
            k = winreg.CreateKey(winreg.HKEY_CURRENT_USER, base)
            winreg.SetValueEx(k, "", 0, winreg.REG_SZ, label)
            winreg.SetValueEx(k, "Icon", 0, winreg.REG_SZ, quoted + ",0")
            ck = winreg.CreateKey(k, "command")
            winreg.SetValueEx(ck, "", 0, winreg.REG_SZ, f'{quoted} --{action} "%1"')
            ck.Close(); k.Close()
        except Exception:
            pass

    try:
        for ext in (".docx", ".xlsx", ".pptx"):
            set_menu(ext, "encrypt", "用Office敏感词工具加密")
            set_menu(ext, "decrypt", "用Office敏感词工具解密")
    except Exception:
        pass


def ensure_shell_integration():
    """打包（exe/MSIX）环境首次启动时安装桌面快捷方式 + 右键菜单；
    每次启动都会校验并修正路径（应对 MSIX 更新后 exe 位置变化）。
    源码运行/非 Windows 环境自动跳过。"""
    if not getattr(sys, "frozen", False) or not sys.platform.startswith("win"):
        return
    try:
        ensure_desktop_shortcut()
        ensure_context_menu()
    except Exception:
        pass


# ==================== 通用工具 ====================

def compute_file_sha256(path: str) -> str:
    """计算文件 SHA-256 指纹，用于词表与文档绑定"""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def restore_doc_name(enc_name: str) -> str:
    """由加密文件名恢复原始文件名: xxx_encrypted.docx -> xxx.docx
    仅在末尾匹配 _encrypted，避免误删文件名中部的同名片段"""
    stem, ext = os.path.splitext(enc_name)
    if stem.endswith("_encrypted"):
        stem = stem[: -len("_encrypted")]
    return stem + ext


# ==================== 依赖检查 ====================

DEPENDENCIES = {
    "docx": "python-docx",
    "openpyxl": "openpyxl",
    "pptx": "python-pptx",
    "cryptography": "cryptography"
}

missing_deps = []
for mod, pkg in DEPENDENCIES.items():
    try:
        __import__(mod)
    except ImportError:
        missing_deps.append(pkg)

if missing_deps:
    print("缺少以下依赖包，请先安装：")
    print(f"  pip install {' '.join(missing_deps)}")
    sys.exit(1)

from docx import Document
from docx.oxml.ns import qn
import openpyxl
from pptx import Presentation
from cryptography.fernet import Fernet, InvalidToken
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

try:
    from openpyxl.cell.rich_text import CellRichText
    HAS_RICH_TEXT = True
except ImportError:
    HAS_RICH_TEXT = False
    CellRichText = None

from decimal import Decimal
from datetime import datetime, date, time, timedelta


# ==================== 加密核心模块 ====================

class CryptoEngine:
    """词表加密引擎
    v2.3: 密钥派生 PBKDF2-HMAC-SHA256(480000次, 16B盐) → AES-256-GCM 认证加密
    文件格式: b"OEV2" + salt(16B) + nonce(12B) + ciphertext+tag
    兼容: 以 16B 随机盐开头的旧版 Fernet 格式仍可解密
    """
    MAGIC = b"OEV2"
    SALT_LEN = 16
    NONCE_LEN = 12
    ITERATIONS = 480000

    @staticmethod
    def derive_key(password: str, salt: bytes) -> bytes:
        kdf = PBKDF2HMAC(
            algorithm=hashes.SHA256(),
            length=32,
            salt=salt,
            iterations=CryptoEngine.ITERATIONS,
        )
        return kdf.derive(password.encode("utf-8"))

    @staticmethod
    def encrypt_vocab(vocab_data: dict, password: str) -> bytes:
        salt = secrets.token_bytes(CryptoEngine.SALT_LEN)
        nonce = secrets.token_bytes(CryptoEngine.NONCE_LEN)
        key = CryptoEngine.derive_key(password, salt)
        payload = json.dumps(vocab_data, ensure_ascii=False).encode("utf-8")
        ciphertext = AESGCM(key).encrypt(nonce, payload, None)
        return CryptoEngine.MAGIC + salt + nonce + ciphertext

    @staticmethod
    def decrypt_vocab(encrypted_data: bytes, password: str) -> dict:
        if len(encrypted_data) < 16:
            raise ValueError("无效的加密文件")

        if encrypted_data[:4] == CryptoEngine.MAGIC:
            # ---- 新版 AES-256-GCM 格式 ----
            if len(encrypted_data) < 4 + CryptoEngine.SALT_LEN + CryptoEngine.NONCE_LEN + 16:
                raise ValueError("无效的加密文件")
            salt = encrypted_data[4: 4 + CryptoEngine.SALT_LEN]
            nonce = encrypted_data[4 + CryptoEngine.SALT_LEN: 4 + CryptoEngine.SALT_LEN + CryptoEngine.NONCE_LEN]
            ciphertext = encrypted_data[4 + CryptoEngine.SALT_LEN + CryptoEngine.NONCE_LEN:]
            key = CryptoEngine.derive_key(password, salt)
            try:
                plain = AESGCM(key).decrypt(nonce, ciphertext, None)
            except Exception:
                raise ValueError("密码错误或文件已损坏")
            return json.loads(plain.decode("utf-8"))
        else:
            # ---- 兼容 v2.2 Fernet 格式: salt(16B) + token ----
            salt = encrypted_data[:16]
            token = encrypted_data[16:]
            key = base64.urlsafe_b64encode(CryptoEngine.derive_key(password, salt))
            f = Fernet(key)
            try:
                decrypted = f.decrypt(token)
                return json.loads(decrypted.decode("utf-8"))
            except InvalidToken:
                raise ValueError("密码错误或文件已损坏")


# ==================== 通用替换函数 ====================

class RunsContainer:
    """文本容器适配器：把任意 run 序列包装成 replace_in_runs 需要的段落接口"""
    def __init__(self, runs):
        self.runs = runs

    @property
    def text(self) -> str:
        return "".join(r.text for r in self.runs)


def replace_each_choice(text: str, old_text: str, choices: list):
    """把 text 中 old_text 的每一次出现分别替换为 choices 中随机一项。
    非重叠扫描，从左到右定位；返回 (新文本, 替换次数)。"""
    parts = []
    count = 0
    i = 0
    while True:
        idx = text.find(old_text, i)
        if idx == -1:
            parts.append(text[i:])
            break
        parts.append(text[i:idx])
        parts.append(secrets.choice(choices))
        i = idx + len(old_text)
        count += 1
    return "".join(parts), count


def replace_in_runs(container, old_text: str, new_text: str, choices: list = None) -> int:
    """在 run 序列中替换文本，支持跨 run 匹配；返回替换次数。
    - choices 为 None：所有出现统一替换为 new_text（标准模式/解密）
    - choices 为列表：每次出现随机选用其中一个字符串（强加密模式）
    """
    if old_text not in container.text:
        return 0

    full_text = ""
    run_map = []
    for run in container.runs:
        start = len(full_text)
        full_text += run.text
        run_map.append((start, start + len(run.text), run))

    # 一次性、非重叠地定位所有出现位置，并为每处选定替换串
    positions = []
    start = 0
    while True:
        idx = full_text.find(old_text, start)
        if idx == -1:
            break
        replacement = secrets.choice(choices) if choices else new_text
        positions.append((idx, replacement))
        start = idx + len(old_text)

    count = 0
    for pos, replacement in reversed(positions):
        count += 1
        end_pos = pos + len(old_text)
        first_idx = None
        last_idx = None
        for i, (s, e, r) in enumerate(run_map):
            if s <= pos < e and first_idx is None:
                first_idx = i
            if s < end_pos <= e:
                last_idx = i
                break
            elif s <= pos < e and e <= end_pos and first_idx is not None:
                last_idx = i

        if first_idx is not None and last_idx is not None:
            merged_text = ""
            for i in range(first_idx, last_idx + 1):
                merged_text += run_map[i][2].text

            offset = pos - run_map[first_idx][0]
            new_merged = merged_text[:offset] + replacement + merged_text[offset + len(old_text):]

            run_map[first_idx][2].text = new_merged
            for i in range(first_idx + 1, last_idx + 1):
                run_map[i][2].text = ""

            new_run_map = []
            new_full = ""
            for i, (s, e, r) in enumerate(run_map):
                if r.text or i == first_idx:
                    ns = len(new_full)
                    new_full += r.text
                    new_run_map.append((ns, ns + len(r.text), r))
            run_map = new_run_map
            full_text = new_full

    return count


# ==================== 文档处理器抽象接口 ====================

class DocumentProcessor(ABC):
    MARKER = "\u200b"

    @abstractmethod
    def load(self, path: str) -> bool:
        pass

    @abstractmethod
    def replace_all(self, mappings: dict, mode: str) -> int:
        pass

    @abstractmethod
    def save(self, path: str) -> bool:
        pass

    @property
    @abstractmethod
    def file_type(self) -> str:
        pass

    def _build_items(self, mappings: dict, mode: str):
        """把映射统一规整为 [(old_text, [replacement, ...]), ...] 有序替换项。
        - 标准模式：mappings 值为字符串（每个敏感词 1 个等长替换串）
        - 强加密模式：mappings 值为字符串列表（每个敏感词 5~10 个变长替换串）
        加密：old=敏感词，new=替换串+零宽标记
        解密：old=替换串+零宽标记，new=敏感词（强加密的每个替换串各占一项）
        长词优先，避免短词误伤。
        """
        items = []
        if mode == "encrypt":
            for word, value in mappings.items():
                if isinstance(value, list):
                    items.append((word, [v + self.MARKER for v in value]))
                else:
                    items.append((word, [value + self.MARKER]))
        else:
            for word, value in mappings.items():
                if isinstance(value, list):
                    for v in value:
                        items.append((v + self.MARKER, [word]))
                else:
                    items.append((value + self.MARKER, [word]))
        items.sort(key=lambda x: len(x[0]), reverse=True)
        return items


class WordProcessor(DocumentProcessor):
    def __init__(self):
        self.doc = None
        self.skipped_warnings = []

    def load(self, path: str) -> bool:
        try:
            self.doc = Document(path)
            return True
        except Exception as e:
            messagebox.showerror("错误", f"无法打开Word文档: {e}")
            return False

    def _iter_text_containers(self):
        """遍历 Word 中所有可替换文本容器：
        正文段落（含表格单元格、文本框 w:txbxContent 内段落）、
        超链接内 run（经 iter_inner_content 合并）、
        已存在的页眉/页脚段落"""
        from docx.text.paragraph import Paragraph

        def build_containers(parent, p_el):
            para = Paragraph(p_el, parent)
            runs = []
            for child in para.iter_inner_content():
                if hasattr(child, "runs"):   # Hyperlink
                    runs.extend(child.runs)
                else:                        # Run
                    runs.append(child)
            if runs:
                yield RunsContainer(runs)

        # 主文档正文及嵌套结构（表格、文本框等）
        for p_el in self.doc.element.body.iter(qn("w:p")):
            yield from build_containers(self.doc, p_el)

        # 页眉 / 页脚（仅处理文档中实际存在的引用，避免创建空页眉改变排版）
        for section in self.doc.sections:
            sectPr = section._sectPr
            if sectPr.find(qn("w:headerReference")) is not None:
                hdr = section.header
                for p_el in hdr._element.iter(qn("w:p")):
                    yield from build_containers(hdr, p_el)
            if sectPr.find(qn("w:footerReference")) is not None:
                ftr = section.footer
                for p_el in ftr._element.iter(qn("w:p")):
                    yield from build_containers(ftr, p_el)

    def replace_all(self, mappings: dict, mode: str) -> int:
        items = self._build_items(mappings, mode)
        total = 0

        for container in self._iter_text_containers():
            for old_text, replacements in items:
                if len(replacements) == 1:
                    total += replace_in_runs(container, old_text, replacements[0])
                else:
                    # 强加密：同词每次出现随机选用不同替换串
                    total += replace_in_runs(container, old_text, None, choices=replacements)

        return total

    def save(self, path: str) -> bool:
        try:
            self.doc.save(path)
            return True
        except Exception as e:
            messagebox.showerror("错误", f"保存Word文档失败: {e}")
            return False

    @property
    def file_type(self) -> str:
        return "Word"


class ExcelProcessor(DocumentProcessor):
    def __init__(self):
        self.wb = None
        self.skipped_warnings = []

    def load(self, path: str) -> bool:
        try:
            self.wb = openpyxl.load_workbook(path, rich_text=True)
            return True
        except Exception as e:
            messagebox.showerror("错误", f"无法打开Excel文档: {e}")
            return False

    def replace_all(self, mappings: dict, mode: str) -> int:
        items = self._build_items(mappings, mode)
        total = 0
        self.skipped_warnings = []

        for sheet_name in self.wb.sheetnames:
            ws = self.wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    for old_text, replacements in items:
                        if len(replacements) == 1:
                            total += self._replace_in_cell(cell, old_text, replacements[0])
                        else:
                            # 强加密：同词每次出现随机选用不同替换串
                            total += self._replace_in_cell(cell, old_text, None,
                                                           choices=replacements)

        return total

    def _replace_in_cell(self, cell, old_text: str, new_text, choices: list = None) -> int:
        value = cell.value
        if value is None:
            return 0

        # ---- 类型保护 ----
        # 布尔 / 日期时间：不参与替换，避免破坏类型
        if isinstance(value, (bool, datetime, date, time, timedelta)):
            return 0

        # 数字（int/float/Decimal）：命中敏感词时尝试保持数值类型
        if isinstance(value, (int, float, Decimal)):
            s = str(value)
            if old_text not in s:
                return 0
            # 强加密替换串为变长字母串，替换后不可能保持数字类型，直接跳过
            if choices:
                self.skipped_warnings.append(
                    f"单元格 {cell.coordinate} 的数字含敏感词 '{old_text}'，"
                    f"强加密替换串无法保持数字类型，已跳过（如需处理请将列格式设为文本）"
                )
                return 0
            new_s = s.replace(old_text, new_text)
            try:
                if isinstance(value, int):
                    cell.value = int(new_s)
                elif isinstance(value, float):
                    cell.value = float(new_s)
                else:
                    cell.value = Decimal(new_s)
                return s.count(old_text)
            except (ValueError, ArithmeticError):
                # 替换后无法保持数字类型（如替换串含字母），跳过并记录提示
                self.skipped_warnings.append(
                    f"单元格 {cell.coordinate} 的数字含敏感词 '{old_text}'，"
                    f"替换后无法保持数字类型，已跳过（如需处理请将列格式设为文本）"
                )
                return 0

        # 公式单元格：跳过，避免破坏公式
        if isinstance(value, str) and value.startswith("="):
            return 0

        is_rich = False
        if HAS_RICH_TEXT and isinstance(value, CellRichText):
            is_rich = True

        cell_val_str = str(value)
        if old_text not in cell_val_str:
            return 0

        pick = choices if choices else [new_text]

        if is_rich:
            full_text = "".join(str(block) for block in value)
            if old_text not in full_text:
                return 0

            count = 0
            found_in_block = False
            for block in value:
                block_text = str(block)
                if old_text in block_text:
                    replaced, n = replace_each_choice(block_text, old_text, pick)
                    if hasattr(block, "text"):
                        block.text = replaced
                    else:
                        idx = value.index(block)
                        value[idx] = replaced
                    count += n
                    found_in_block = True

            if not found_in_block:
                # 敏感词跨越多个格式块：合并为纯文本（内容正确，格式合并）
                new_full, count = replace_each_choice(full_text, old_text, pick)
                if len(value) > 0:
                    first_block = value[0]
                    if hasattr(first_block, "text"):
                        first_block.text = new_full
                    else:
                        value[0] = new_full
                    for i in range(1, len(value)):
                        if hasattr(value[i], "text"):
                            value[i].text = ""
                        else:
                            value[i] = ""

            return count
        else:
            new_val, count = replace_each_choice(cell_val_str, old_text, pick)
            cell.value = new_val
            return count

    def save(self, path: str) -> bool:
        try:
            self.wb.save(path)
            return True
        except Exception as e:
            messagebox.showerror("错误", f"保存Excel文档失败: {e}")
            return False

    @property
    def file_type(self) -> str:
        return "Excel"


class PptProcessor(DocumentProcessor):
    def __init__(self):
        self.prs = None
        self.skipped_warnings = []

    def load(self, path: str) -> bool:
        try:
            self.prs = Presentation(path)
            return True
        except Exception as e:
            messagebox.showerror("错误", f"无法打开PPT文档: {e}")
            return False

    def replace_all(self, mappings: dict, mode: str) -> int:
        items = self._build_items(mappings, mode)
        total = 0

        def replace_para(para):
            nonlocal total
            for old_text, replacements in items:
                if len(replacements) == 1:
                    total += replace_in_runs(para, old_text, replacements[0])
                else:
                    # 强加密：同词每次出现随机选用不同替换串
                    total += replace_in_runs(para, old_text, None, choices=replacements)

        for slide in self.prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        replace_para(para)

                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            for para in cell.text_frame.paragraphs:
                                replace_para(para)

            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                for para in notes_slide.notes_text_frame.paragraphs:
                    replace_para(para)

        return total

    def save(self, path: str) -> bool:
        try:
            self.prs.save(path)
            return True
        except Exception as e:
            messagebox.showerror("错误", f"保存PPT文档失败: {e}")
            return False

    @property
    def file_type(self) -> str:
        return "PowerPoint"


def get_processor(file_path: str) -> DocumentProcessor:
    ext = Path(file_path).suffix.lower()
    if ext == ".docx":
        return WordProcessor()
    elif ext == ".xlsx":
        return ExcelProcessor()
    elif ext in (".pptx", ".ppsx"):
        return PptProcessor()
    else:
        raise ValueError(f"不支持的文件格式: {ext}")


class SubstitutionEngine:
    CHARS = string.ascii_letters
    STRONG_MIN_VARIANTS = 5
    STRONG_MAX_VARIANTS = 10

    @classmethod
    def generate(cls, length: int) -> str:
        return "".join(secrets.choice(cls.CHARS) for _ in range(length))

    @classmethod
    def create_mapping(cls, words: list) -> dict:
        """标准模式：每个敏感词 -> 1 个等长随机英文字符串 {word: sub}"""
        mapping = {}
        used = set()
        for word in words:
            word = word.strip()
            if not word:
                continue
            while True:
                sub = cls.generate(len(word))
                if sub not in used and sub != word:
                    used.add(sub)
                    break
            mapping[word] = sub
        return mapping

    @classmethod
    def create_strong_mapping(cls, words: list) -> dict:
        """强加密模式：每个敏感词 -> 5~10 个随机英文字符串 {word: [sub,...]}
        - 每个字符串长度在 [L, 2L] 之间随机（L = 敏感词长度）
        - 同一敏感词在文档中的每次出现，随机选用其中一个字符串替换
        """
        mapping = {}
        used = set()
        word_set = {w.strip() for w in words if w.strip()}
        for word in words:
            word = word.strip()
            if not word:
                continue
            length = len(word)
            variant_count = cls.STRONG_MIN_VARIANTS + \
                secrets.randbelow(cls.STRONG_MAX_VARIANTS - cls.STRONG_MIN_VARIANTS + 1)
            variants = []
            while len(variants) < variant_count:
                # 长度随机落在 [L, 2L]
                sub_len = length + secrets.randbelow(length + 1)
                sub = cls.generate(sub_len)
                if sub in used or sub in word_set:
                    continue
                used.add(sub)
                variants.append(sub)
            mapping[word] = variants
        return mapping


# ==================== 后台任务线程 ====================

class TaskThread:
    """后台任务执行器：worker 线程跑耗时操作，日志经队列回传主线程，
    完成或异常后回调 on_done。所有 GUI 操作只发生在主线程。"""

    def __init__(self, root, on_message, on_done):
        self.root = root
        self.on_message = on_message
        self.on_done = on_done
        self.q = queue.Queue()
        self.running = False

    def start(self, func, *args):
        if self.running:
            return
        self.running = True
        t = threading.Thread(target=self._run, args=(func, args), daemon=True)
        t.start()
        self._poll()

    def _run(self, func, args):
        try:
            func(*args)
            self.q.put(("done", None))
        except Exception as e:
            import traceback
            self.q.put(("done", (e, traceback.format_exc())))

    def _poll(self):
        try:
            while True:
                kind, payload = self.q.get_nowait()
                if kind == "done":
                    self.running = False
                    self.on_done(payload)
                    return
                elif kind == "log":
                    self.on_message(payload)
                else:
                    # 其它结构化消息（如批量统计 result）暂存，由 on_done 读取
                    if not hasattr(self, "results"):
                        self.results = []
                    self.results.append((kind, payload))
        except queue.Empty:
            pass
        self.root.after(80, self._poll)

    def is_running(self):
        return self.running


def check_overwrite_conflicts(paths, single: bool) -> str:
    """主线程中检查目标文件是否已存在，返回处理策略：
    'ok' 无冲突；'overwrite' 全部覆盖；'skip' 跳过已存在；'cancel' 中止"""
    existing = [p for p in paths if os.path.exists(p)]
    if not existing:
        return "ok"
    if single:
        ok = messagebox.askyesno(
            "文件已存在",
            f"目标文件已存在：\n{existing[0]}\n\n是否覆盖？")
        return "overwrite" if ok else "cancel"
    r = messagebox.askyesnocancel(
        "文件已存在",
        f"{len(existing)} 个目标文件已存在。\n\n"
        f"是(Y) = 全部覆盖\n否(N) = 跳过已存在的文件\n取消 = 中止操作")
    if r is True:
        return "overwrite"
    if r is False:
        return "skip"
    return "cancel"


# ==================== GUI 界面 ====================

class VocabManagerFrame(ttk.LabelFrame):
    def __init__(self, parent, app):
        super().__init__(parent, text=" 敏感词表管理 ", padding=6)
        self.app = app
        self.words = []
        self.build_ui()

    def build_ui(self):
        input_frame = ttk.Frame(self)
        input_frame.pack(fill=tk.X, pady=5)

        ttk.Label(input_frame, text="敏感词:").pack(side=tk.LEFT, padx=5)
        self.word_entry = ttk.Entry(input_frame, width=20)
        self.word_entry.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        self.word_entry.bind("<Return>", lambda e: self.add_word())

        ttk.Button(input_frame, text="添加", command=self.add_word, width=6).pack(side=tk.LEFT, padx=2)
        ttk.Button(input_frame, text="删除", command=self.delete_word, width=6).pack(side=tk.LEFT, padx=2)
        ttk.Button(input_frame, text="清空", command=self.clear_words, width=6).pack(side=tk.LEFT, padx=2)

        list_frame = ttk.Frame(self)
        list_frame.pack(fill=tk.BOTH, expand=True, pady=5)

        scrollbar = ttk.Scrollbar(list_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        self.word_listbox = tk.Listbox(list_frame, yscrollcommand=scrollbar.set,
                                       height=5, font=("Consolas", 10),
                                       bg="#FFFFFF", fg="#1F1F1F",
                                       selectbackground="#0078D4", selectforeground="#FFFFFF",
                                       highlightthickness=1, highlightbackground="#D1D1D1",
                                       highlightcolor="#0078D4", relief="flat")
        self.word_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.config(command=self.word_listbox.yview)

        file_frame = ttk.Frame(self)
        file_frame.pack(fill=tk.X, pady=5)

        ttk.Button(file_frame, text="加载词表", command=self.load_vocab_file).pack(side=tk.LEFT, padx=3, fill=tk.X, expand=True)
        ttk.Button(file_frame, text="保存词表", command=self.save_vocab_file).pack(side=tk.LEFT, padx=3, fill=tk.X, expand=True)
        ttk.Button(file_frame, text="示例词表", command=self.load_demo).pack(side=tk.LEFT, padx=3, fill=tk.X, expand=True)

        self.status_label = ttk.Label(self, text="当前词表: 0 个词汇", foreground="gray")
        self.status_label.pack(anchor=tk.W, pady=(5, 0))

    def add_word(self):
        word = self.word_entry.get().strip()
        if not word:
            messagebox.showwarning("提示", "请输入敏感词")
            return
        if word in self.words:
            messagebox.showwarning("提示", "该词汇已存在")
            return
        self.words.append(word)
        self.word_listbox.insert(tk.END, word)
        self.word_entry.delete(0, tk.END)
        self.update_status()

    def delete_word(self):
        selection = self.word_listbox.curselection()
        if not selection:
            messagebox.showwarning("提示", "请先选择要删除的词汇")
            return
        idx = selection[0]
        del self.words[idx]
        self.word_listbox.delete(idx)
        self.update_status()

    def clear_words(self):
        if not self.words:
            return
        if messagebox.askyesno("确认", "确定清空所有词汇？"):
            self.words.clear()
            self.word_listbox.delete(0, tk.END)
            self.update_status()

    def update_status(self):
        self.status_label.config(text=f"当前词表: {len(self.words)} 个词汇")

    def get_words(self):
        return self.words.copy()

    def set_words(self, words):
        self.words = list(words)
        self.word_listbox.delete(0, tk.END)
        for w in self.words:
            self.word_listbox.insert(tk.END, w)
        self.update_status()

    def load_vocab_file(self):
        path = filedialog.askopenfilename(
            title="选择词表文件",
            initialdir=get_word_dir(),
            filetypes=[("词表文件", "*.json *.vocab"), ("所有文件", "*.*")]
        )
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            words = data.get("words", [])
            self.set_words(words)
            messagebox.showinfo("成功", f"已加载 {len(words)} 个词汇")
        except Exception as e:
            messagebox.showerror("错误", f"加载失败: {e}")

    def save_vocab_file(self):
        if not self.words:
            messagebox.showwarning("提示", "词表为空")
            return
        path = filedialog.asksaveasfilename(
            title="保存词表",
            initialdir=get_word_dir(),
            defaultextension=".json",
            filetypes=[("JSON文件", "*.json"), ("词表文件", "*.vocab"), ("所有文件", "*.*")]
        )
        if not path:
            return
        try:
            data = {"words": self.words, "version": "1.0"}
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            messagebox.showinfo("成功", f"已保存到: {path}")
        except Exception as e:
            messagebox.showerror("错误", f"保存失败: {e}")

    def load_demo(self):
        demo = ["机密", "绝密", "内部资料", "薪酬", "身份证号", "银行卡号",
                "密码", "密钥", "商业机密", "客户名单", "合同金额", "谈判底线"]
        self.set_words(demo)
        messagebox.showinfo("提示", "已加载示例词表")


class EncryptFrame(ttk.LabelFrame):
    def __init__(self, parent, app):
        super().__init__(parent, text=" 文档加密 ", padding=6)
        self.app = app
        self.doc_path = None
        self.processor = None
        self.task = None
        self.build_ui()

    def build_ui(self):
        doc_frame = ttk.Frame(self)
        doc_frame.pack(fill=tk.X, pady=3)
        ttk.Label(doc_frame, text="文档:").pack(side=tk.LEFT, padx=5)
        self.doc_label = ttk.Label(doc_frame, text="未选择", foreground="gray")
        self.doc_label.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        ttk.Button(doc_frame, text="浏览...", command=self.select_doc, width=8).pack(side=tk.LEFT, padx=5)

        pwd_frame = ttk.Frame(self)
        pwd_frame.pack(fill=tk.X, pady=3)
        ttk.Label(pwd_frame, text="密码:").pack(side=tk.LEFT, padx=5)
        self.pwd_entry = ttk.Entry(pwd_frame, show="*", width=18)
        self.pwd_entry.pack(side=tk.LEFT, padx=5)
        self.show_pwd_var = tk.BooleanVar()
        ttk.Checkbutton(pwd_frame, text="显示", variable=self.show_pwd_var,
                       command=self.toggle_pwd).pack(side=tk.LEFT, padx=5)
        self.strong_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(pwd_frame, text="强加密", variable=self.strong_var).pack(side=tk.LEFT, padx=5)

        btn_frame = ttk.Frame(self)
        btn_frame.pack(fill=tk.X, pady=5)
        self.encrypt_btn = ttk.Button(btn_frame, text="一键加密", command=self.encrypt, width=18, style="Accent.TButton")
        self.encrypt_btn.pack(pady=2)

        ttk.Label(self, text="日志:").pack(anchor=tk.W, pady=(3, 0))
        self.log_text = scrolledtext.ScrolledText(self, height=3, state=tk.DISABLED, font=("Consolas", 9), bg="#FFFFFF", fg="#1F1F1F", insertbackground="#1F1F1F", relief="flat", highlightthickness=1, highlightbackground="#D1D1D1", highlightcolor="#0078D4")
        self.log_text.pack(fill=tk.BOTH, expand=True, pady=3)

    def toggle_pwd(self):
        self.pwd_entry.config(show="" if self.show_pwd_var.get() else "*")

    def select_doc(self):
        path = filedialog.askopenfilename(
            title="选择Office文档",
            filetypes=[
                ("Office文档", "*.docx *.xlsx *.pptx"),
                ("Word", "*.docx"),
                ("Excel", "*.xlsx"),
                ("PowerPoint", "*.pptx"),
                ("所有文件", "*.*")
            ]
        )
        if path:
            self.set_doc(path)

    def set_doc(self, path):
        """外部（右键菜单/命令行）指定文档"""
        try:
            self.processor = get_processor(path)
            self.doc_path = path
            name = Path(path).name
            self.doc_label.config(
                text=f"[{self.processor.file_type}] {name if len(name) < 25 else name[:22] + '...'}",
                foreground="black"
            )
        except ValueError as e:
            messagebox.showerror("错误", str(e))

    def log(self, msg):
        self.log_text.config(state=tk.NORMAL)
        self.log_text.insert(tk.END, msg + "\n")
        self.log_text.see(tk.END)
        self.log_text.config(state=tk.DISABLED)
        self.app.update_idletasks()

    def encrypt(self):
        words = self.app.vocab_frame.get_words()
        if not words:
            messagebox.showwarning("提示", "请先添加敏感词")
            return
        if not self.doc_path or not self.processor:
            messagebox.showwarning("提示", "请选择Office文档")
            return
        password = self.pwd_entry.get()
        if not password:
            messagebox.showwarning("提示", "请设置解密密码")
            return
        if len(password) < 4:
            messagebox.showwarning("提示", "密码长度至少4位")
            return
        if self.task and self.task.is_running():
            messagebox.showwarning("提示", "正在处理中，请稍候...")
            return

        save_dir = filedialog.askdirectory(title="选择加密文件保存位置")
        if not save_dir:
            return

        base_name = Path(self.doc_path).stem
        ext = Path(self.doc_path).suffix
        enc_doc_path = os.path.join(save_dir, f"{base_name}_encrypted{ext}")
        enc_vocab_path = os.path.join(save_dir, f"{base_name}_vocab.enc")

        # 主线程预检覆盖
        strategy = check_overwrite_conflicts([enc_doc_path, enc_vocab_path], single=True)
        if strategy == "cancel":
            return

        strong = self.strong_var.get()  # 主线程读取 tk 变量

        self.encrypt_btn.config(state=tk.DISABLED)
        self.log("=" * 40)
        self.log("开始加密（强加密模式）..." if strong else "开始加密...")

        self.task = TaskThread(
            self.app.root,
            on_message=self.log,
            on_done=lambda err: self._encrypt_done(err, enc_doc_path, enc_vocab_path)
        )
        self.task.start(self._encrypt_worker, words, password, save_dir,
                        enc_doc_path, enc_vocab_path, strong)

    def _encrypt_worker(self, words, password, save_dir, enc_doc_path, enc_vocab_path,
                        strong=False):
        """后台线程执行：生成映射 → 替换文档 → 保存 → 计算加密文档指纹 → 加密词表"""
        log = lambda m: self.task.q.put(("log", m))

        log("正在生成替换映射...")
        if strong:
            log("强加密：每个敏感词生成 5~10 个、长度 L~2L 的随机替换串，逐处随机选用")
            mappings = SubstitutionEngine.create_strong_mapping(words)
        else:
            mappings = SubstitutionEngine.create_mapping(words)
        log("替换映射表:")
        for w, m in mappings.items():
            if isinstance(m, list):
                log(f"  {w} -> [{len(m)} 个] " + ", ".join(m))
            else:
                log(f"  {w} -> {m}")

        log(f"正在处理{self.processor.file_type}文档: {Path(self.doc_path).name}...")
        if not self.processor.load(self.doc_path):
            raise RuntimeError("文档加载失败")
        count = self.processor.replace_all(mappings, mode="encrypt")

        if not self.processor.save(enc_doc_path):
            raise RuntimeError("加密文档保存失败")

        # 文档指纹绑定：记录加密后文档的 SHA-256
        log("正在计算文档指纹...")
        doc_hash = compute_file_sha256(enc_doc_path)
        vocab_data = {"version": "2.4", "words": words,
                      "mappings": mappings, "doc_hash": doc_hash,
                      "strong": strong}

        log("正在加密词表 (AES-256-GCM)...")
        encrypted_vocab = CryptoEngine.encrypt_vocab(vocab_data, password)
        with open(enc_vocab_path, "wb") as f:
            f.write(encrypted_vocab)

        log(f"加密完成！共替换 {count} 处")
        log(f"加密文档: {enc_doc_path}")
        log(f"加密词表: {enc_vocab_path}")

    def _encrypt_done(self, err, enc_doc_path, enc_vocab_path):
        self.encrypt_btn.config(state=tk.NORMAL)
        if err is None:
            self.log("全部完成。")
            messagebox.showinfo("成功",
                f"加密完成！\n\n"
                f"文档类型: {self.processor.file_type}\n"
                f"加密文档: {enc_doc_path}\n"
                f"加密词表: {enc_vocab_path}\n\n"
                f"请妥善保管解密密码！")
        else:
            exc, tb = err
            self.log(f"加密失败: {exc}")
            messagebox.showerror("错误", f"加密失败: {exc}")


class DecryptFrame(ttk.LabelFrame):
    def __init__(self, parent, app):
        super().__init__(parent, text=" 文档解密 ", padding=6)
        self.app = app
        self.doc_path = None
        self.vocab_path = None
        self.processor = None
        self.task = None
        self.build_ui()

    def build_ui(self):
        doc_frame = ttk.Frame(self)
        doc_frame.pack(fill=tk.X, pady=3)
        ttk.Label(doc_frame, text="文档:").pack(side=tk.LEFT, padx=5)
        self.doc_label = ttk.Label(doc_frame, text="未选择", foreground="gray")
        self.doc_label.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        ttk.Button(doc_frame, text="浏览...", command=self.select_doc, width=8).pack(side=tk.LEFT, padx=5)

        vocab_frame = ttk.Frame(self)
        vocab_frame.pack(fill=tk.X, pady=3)
        ttk.Label(vocab_frame, text="词表:").pack(side=tk.LEFT, padx=5)
        self.vocab_label = ttk.Label(vocab_frame, text="未选择", foreground="gray")
        self.vocab_label.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        ttk.Button(vocab_frame, text="浏览...", command=self.select_vocab, width=8).pack(side=tk.LEFT, padx=5)

        pwd_frame = ttk.Frame(self)
        pwd_frame.pack(fill=tk.X, pady=3)
        ttk.Label(pwd_frame, text="密码:").pack(side=tk.LEFT, padx=5)
        self.pwd_entry = ttk.Entry(pwd_frame, show="*", width=18)
        self.pwd_entry.pack(side=tk.LEFT, padx=5)
        self.show_pwd_var = tk.BooleanVar()
        ttk.Checkbutton(pwd_frame, text="显示", variable=self.show_pwd_var,
                       command=self.toggle_pwd).pack(side=tk.LEFT, padx=5)

        btn_frame = ttk.Frame(self)
        btn_frame.pack(fill=tk.X, pady=5)
        self.decrypt_btn = ttk.Button(btn_frame, text="一键解密", command=self.decrypt, width=18, style="Accent.TButton")
        self.decrypt_btn.pack(pady=2)

        ttk.Label(self, text="日志:").pack(anchor=tk.W, pady=(3, 0))
        self.log_text = scrolledtext.ScrolledText(self, height=3, state=tk.DISABLED, font=("Consolas", 9), bg="#FFFFFF", fg="#1F1F1F", insertbackground="#1F1F1F", relief="flat", highlightthickness=1, highlightbackground="#D1D1D1", highlightcolor="#0078D4")
        self.log_text.pack(fill=tk.BOTH, expand=True, pady=3)

    def toggle_pwd(self):
        self.pwd_entry.config(show="" if self.show_pwd_var.get() else "*")

    def select_doc(self):
        path = filedialog.askopenfilename(
            title="选择加密后的Office文档",
            filetypes=[
                ("Office文档", "*.docx *.xlsx *.pptx"),
                ("Word", "*.docx"),
                ("Excel", "*.xlsx"),
                ("PowerPoint", "*.pptx"),
                ("所有文件", "*.*")
            ]
        )
        if path:
            self.set_doc(path)

    def set_doc(self, path):
        """外部（右键菜单/命令行）指定文档"""
        try:
            self.processor = get_processor(path)
            self.doc_path = path
            name = Path(path).name
            self.doc_label.config(
                text=f"[{self.processor.file_type}] {name if len(name) < 25 else name[:22] + '...'}",
                foreground="black"
            )
        except ValueError as e:
            messagebox.showerror("错误", str(e))

    def select_vocab(self):
        path = filedialog.askopenfilename(
            title="选择加密词表",
            filetypes=[("加密词表", "*.enc"), ("所有文件", "*.*")]
        )
        if not path:
            return
        self.vocab_path = path
        name = Path(path).name
        self.vocab_label.config(text=name if len(name) < 30 else name[:27] + "...", foreground="black")

    def log(self, msg):
        self.log_text.config(state=tk.NORMAL)
        self.log_text.insert(tk.END, msg + "\n")
        self.log_text.see(tk.END)
        self.log_text.config(state=tk.DISABLED)
        self.app.update_idletasks()

    def decrypt(self):
        if not self.doc_path or not self.processor:
            messagebox.showwarning("提示", "请选择加密文档")
            return
        if not self.vocab_path:
            messagebox.showwarning("提示", "请选择加密词表")
            return
        password = self.pwd_entry.get()
        if not password:
            messagebox.showwarning("提示", "请输入解密密码")
            return
        if self.task and self.task.is_running():
            messagebox.showwarning("提示", "正在处理中，请稍候...")
            return

        save_dir = filedialog.askdirectory(title="选择解密文件保存位置")
        if not save_dir:
            return

        dec_name = restore_doc_name(Path(self.doc_path).name)
        dec_doc_path = os.path.join(save_dir, dec_name)
        dec_vocab_path = os.path.join(save_dir, Path(dec_name).stem + "_vocab.json")

        strategy = check_overwrite_conflicts([dec_doc_path, dec_vocab_path], single=True)
        if strategy == "cancel":
            return

        self.decrypt_btn.config(state=tk.DISABLED)
        self.log("=" * 40)
        self.log("开始解密...")

        self.task = TaskThread(
            self.app.root,
            on_message=self.log,
            on_done=lambda err: self._decrypt_done(err, dec_doc_path, dec_vocab_path)
        )
        self.task.start(self._decrypt_worker, password, save_dir, dec_doc_path, dec_vocab_path)

    def _decrypt_worker(self, password, save_dir, dec_doc_path, dec_vocab_path):
        log = lambda m: self.task.q.put(("log", m))

        log("正在解密词表...")
        with open(self.vocab_path, "rb") as f:
            encrypted_vocab = f.read()

        vocab_data = CryptoEngine.decrypt_vocab(encrypted_vocab, password)
        words = vocab_data.get("words", [])
        mappings = vocab_data.get("mappings", {})

        # 词表-文档绑定校验
        log("正在校验文档指纹...")
        doc_hash = compute_file_sha256(self.doc_path)
        bound_hash = vocab_data.get("doc_hash")
        if bound_hash:
            if bound_hash != doc_hash:
                raise ValueError("文档与词表不匹配：该词表不是为当前文档生成的")
            log("文档指纹校验通过")
        else:
            log("提示: 该词表由旧版本生成，未包含文档指纹，跳过校验")

        log(f"词表解密成功，包含 {len(words)} 个词汇")
        log("词表类型: " + ("强加密（多变体变长替换）" if vocab_data.get("strong") else "标准加密"))
        log("替换映射表:")
        for w, m in mappings.items():
            if isinstance(m, list):
                for v in m:
                    log(f"  {v} -> {w}")
            else:
                log(f"  {m} -> {w}")

        log(f"正在处理{self.processor.file_type}文档: {Path(self.doc_path).name}...")
        if not self.processor.load(self.doc_path):
            raise RuntimeError("文档加载失败")
        count = self.processor.replace_all(mappings, mode="decrypt")

        if not self.processor.save(dec_doc_path):
            raise RuntimeError("解密文档保存失败")

        with open(dec_vocab_path, "w", encoding="utf-8") as f:
            json.dump(vocab_data, f, ensure_ascii=False, indent=2)

        log(f"解密完成！共恢复 {count} 处")
        log(f"恢复文档: {dec_doc_path}")
        log(f"恢复词表: {dec_vocab_path}")

    def _decrypt_done(self, err, dec_doc_path, dec_vocab_path):
        self.decrypt_btn.config(state=tk.NORMAL)
        if err is None:
            self.log("全部完成。")
            messagebox.showinfo("成功",
                f"解密完成！\n\n"
                f"文档类型: {self.processor.file_type}\n"
                f"恢复文档: {dec_doc_path}\n"
                f"恢复词表: {dec_vocab_path}")
        else:
            exc, tb = err
            self.log(f"解密失败: {exc}")
            messagebox.showerror("错误", f"解密失败: {exc}")


class BatchFrame(ttk.LabelFrame):
    """批量处理：对一个目录中的文档批量加密 / 批量解密"""

    OFFICE_EXTS = (".docx", ".xlsx", ".pptx")

    def __init__(self, parent, app):
        super().__init__(parent, text=" 批量处理 ", padding=8)
        self.app = app
        self.task = None
        self.build_ui()

    def build_ui(self):
        # --- 密码与选项行（批量加密/解密共用） ---
        pwd_row = ttk.Frame(self)
        pwd_row.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(pwd_row, text="密码:").pack(side=tk.LEFT, padx=(2, 4))
        self.pwd_entry = ttk.Entry(pwd_row, show="*", width=18)
        self.pwd_entry.pack(side=tk.LEFT, padx=4)
        self.show_pwd_var = tk.BooleanVar()
        ttk.Checkbutton(pwd_row, text="显示", variable=self.show_pwd_var,
                        command=self.toggle_pwd).pack(side=tk.LEFT, padx=2)
        self.independent_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(pwd_row, text="每文件独立词表", variable=self.independent_var).pack(side=tk.LEFT, padx=16)
        self.strong_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(pwd_row, text="强加密（批量加密生效）", variable=self.strong_var).pack(side=tk.LEFT, padx=8)
        ttk.Label(pwd_row, text="批量加密与批量解密共用此密码",
                  foreground="gray").pack(side=tk.LEFT, padx=8)

        # --- 批量加密区 ---
        enc_lf = ttk.LabelFrame(self, text=" 批量加密 ", padding=6)
        enc_lf.pack(fill=tk.X, pady=(0, 8))
        enc_row = ttk.Frame(enc_lf)
        enc_row.pack(fill=tk.X, pady=3)
        ttk.Label(enc_row, text="加密目录:").pack(side=tk.LEFT, padx=(2, 4))
        self.enc_src_var = tk.StringVar()
        ttk.Entry(enc_row, textvariable=self.enc_src_var, width=30).pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(enc_row, text="浏览...", width=8, command=self.pick_enc_src).pack(side=tk.LEFT, padx=4)
        ttk.Label(enc_row, text="输出目录:").pack(side=tk.LEFT, padx=(14, 4))
        self.enc_out_var = tk.StringVar()
        ttk.Entry(enc_row, textvariable=self.enc_out_var, width=20).pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(enc_row, text="浏览...", width=8, command=self.pick_enc_out).pack(side=tk.LEFT, padx=4)
        ttk.Label(enc_row, text="留空=源目录", foreground="gray").pack(side=tk.LEFT, padx=(8, 0))
        self.batch_enc_btn = ttk.Button(enc_lf, text="批量加密", command=self.batch_encrypt, width=14, style="Accent.TButton")
        self.batch_enc_btn.pack(pady=4)

        # --- 批量解密区 ---
        dec_lf = ttk.LabelFrame(self, text=" 批量解密 ", padding=6)
        dec_lf.pack(fill=tk.X, pady=(0, 8))
        dec_row = ttk.Frame(dec_lf)
        dec_row.pack(fill=tk.X, pady=3)
        ttk.Label(dec_row, text="解密目录:").pack(side=tk.LEFT, padx=(2, 4))
        self.dec_src_var = tk.StringVar()
        ttk.Entry(dec_row, textvariable=self.dec_src_var, width=30).pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(dec_row, text="浏览...", width=8, command=self.pick_dec_src).pack(side=tk.LEFT, padx=4)
        ttk.Label(dec_row, text="词表目录:").pack(side=tk.LEFT, padx=(14, 4))
        self.dec_vocab_var = tk.StringVar()
        ttk.Entry(dec_row, textvariable=self.dec_vocab_var, width=20).pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(dec_row, text="浏览...", width=8, command=self.pick_dec_vocab).pack(side=tk.LEFT, padx=4)

        dec_row2 = ttk.Frame(dec_lf)
        dec_row2.pack(fill=tk.X, pady=3)
        ttk.Label(dec_row2, text="输出目录:").pack(side=tk.LEFT, padx=(2, 4))
        self.dec_out_var = tk.StringVar()
        ttk.Entry(dec_row2, textvariable=self.dec_out_var, width=30).pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(dec_row2, text="浏览...", width=8, command=self.pick_dec_out).pack(side=tk.LEFT, padx=4)
        ttk.Label(dec_row2, text="留空=源目录", foreground="gray").pack(side=tk.LEFT, padx=(8, 0))
        self.batch_dec_btn = ttk.Button(dec_lf, text="批量解密", command=self.batch_decrypt, width=14, style="Accent.TButton")
        self.batch_dec_btn.pack(pady=4)

        # --- 日志 ---
        ttk.Label(self, text="日志:").pack(anchor=tk.W, pady=(2, 0))
        self.log_text = scrolledtext.ScrolledText(self, height=6, state=tk.DISABLED, font=("Consolas", 9), bg="#FFFFFF", fg="#1F1F1F", insertbackground="#1F1F1F", relief="flat", highlightthickness=1, highlightbackground="#D1D1D1", highlightcolor="#0078D4")
        self.log_text.pack(fill=tk.BOTH, expand=True, pady=3)

    # ---------- 目录选择 ----------
    def pick_enc_src(self):
        d = filedialog.askdirectory(title="选择待加密文档所在目录")
        if d:
            self.enc_src_var.set(d)

    def pick_enc_out(self):
        d = filedialog.askdirectory(title="选择加密输出目录")
        if d:
            self.enc_out_var.set(d)

    def pick_dec_src(self):
        d = filedialog.askdirectory(title="选择加密文档所在目录")
        if d:
            self.dec_src_var.set(d)

    def pick_dec_vocab(self):
        d = filedialog.askdirectory(title="选择词表所在目录")
        if d:
            self.dec_vocab_var.set(d)

    def pick_dec_out(self):
        d = filedialog.askdirectory(title="选择解密输出目录")
        if d:
            self.dec_out_var.set(d)

    def toggle_pwd(self):
        self.pwd_entry.config(show="" if self.show_pwd_var.get() else "*")

    # ---------- 工具 ----------
    def log(self, msg):
        self.log_text.config(state=tk.NORMAL)
        self.log_text.insert(tk.END, msg + "\n")
        self.log_text.see(tk.END)
        self.log_text.config(state=tk.DISABLED)
        self.app.update_idletasks()

    @staticmethod
    def list_office_files(directory: str):
        """列出目录中支持处理的 Office 文件（一级目录，不递归）"""
        files = []
        try:
            for name in sorted(os.listdir(directory)):
                full = os.path.join(directory, name)
                if os.path.isfile(full) and Path(name).suffix.lower() in BatchFrame.OFFICE_EXTS:
                    files.append(full)
        except OSError as e:
            raise ValueError(f"无法读取目录: {e}")
        return files

    # ---------- 批量加密 ----------
    def batch_encrypt(self):
        words = self.app.vocab_frame.get_words()
        if not words:
            messagebox.showwarning("提示", "请先添加敏感词")
            return
        src_dir = self.enc_src_var.get().strip()
        if not src_dir or not os.path.isdir(src_dir):
            messagebox.showwarning("提示", "请选择有效的加密源目录")
            return
        out_dir = self.enc_out_var.get().strip() or src_dir
        password = self.pwd_entry.get()
        if not password:
            messagebox.showwarning("提示", "请设置解密密码")
            return
        if len(password) < 4:
            messagebox.showwarning("提示", "密码长度至少4位")
            return
        if self.task and self.task.is_running():
            messagebox.showwarning("提示", "正在处理中，请稍候...")
            return

        files = self.list_office_files(src_dir)
        if not files:
            messagebox.showinfo("提示", "目录中没有可处理的 Office 文档")
            return

        independent = self.independent_var.get()  # 主线程读取 tk 变量
        strong = self.strong_var.get()

        # 预计算目标路径，主线程检查覆盖
        targets = []
        for f in files:
            base = Path(f).stem
            ext = Path(f).suffix
            targets.append(os.path.join(out_dir, f"{base}_encrypted{ext}"))
        if not independent:
            targets.append(os.path.join(out_dir, "_batch_vocab.enc"))

        strategy = check_overwrite_conflicts(targets, single=False)
        if strategy == "cancel":
            return

        self.batch_enc_btn.config(state=tk.DISABLED)
        self.batch_dec_btn.config(state=tk.DISABLED)
        self.log("=" * 40)
        self.log(f"批量加密开始：{len(files)} 个文件")

        self.task = TaskThread(
            self.app.root,
            on_message=self.log,
            on_done=lambda err: self._batch_done(err, "encrypt", files, out_dir)
        )
        self.task.start(self._batch_encrypt_worker, words, password, files, out_dir,
                        strategy, independent, strong)

    def _batch_encrypt_worker(self, words, password, files, out_dir, strategy,
                              independent, strong=False):
        log = lambda m: self.task.q.put(("log", m))
        os.makedirs(out_dir, exist_ok=True)
        total = len(files)
        ok_count = 0
        skip_count = 0
        fail_list = []

        if independent:
            # 每文件独立映射 + 独立词表
            for i, f in enumerate(files, 1):
                log(f"[{i}/{total}] 处理: {Path(f).name}")
                try:
                    processor = get_processor(f)
                    if not processor.load(f):
                        raise RuntimeError("文档加载失败")
                    if strong:
                        mappings = SubstitutionEngine.create_strong_mapping(words)
                    else:
                        mappings = SubstitutionEngine.create_mapping(words)
                    count = processor.replace_all(mappings, mode="encrypt")

                    ext = Path(f).suffix
                    enc_doc = os.path.join(out_dir, f"{Path(f).stem}_encrypted{ext}")
                    if strategy == "skip" and os.path.exists(enc_doc):
                        log("  目标已存在，跳过")
                        skip_count += 1
                        continue
                    if not processor.save(enc_doc):
                        raise RuntimeError("加密文档保存失败")

                    doc_hash = compute_file_sha256(enc_doc)
                    vocab_data = {"version": "2.4", "words": words,
                                  "mappings": mappings, "doc_hash": doc_hash,
                                  "strong": strong}
                    enc_vocab = os.path.join(out_dir, f"{Path(f).stem}_vocab.enc")
                    with open(enc_vocab, "wb") as wf:
                        wf.write(CryptoEngine.encrypt_vocab(vocab_data, password))
                    ok_count += 1
                    log(f"  ✓ 替换 {count} 处 -> {Path(enc_doc).name}")
                except Exception as e:
                    fail_list.append(f"{Path(f).name}: {e}")
                    log(f"  ✗ 失败: {e}")
        else:
            # 共享映射 + 共享词表
            log("正在生成共享替换映射...")
            if strong:
                log("强加密：每个敏感词 5~10 个、长度 L~2L 的随机替换串，逐处随机选用")
                mappings = SubstitutionEngine.create_strong_mapping(words)
            else:
                mappings = SubstitutionEngine.create_mapping(words)
            for w, m in mappings.items():
                if isinstance(m, list):
                    log(f"  {w} -> [{len(m)} 个] " + ", ".join(m))
                else:
                    log(f"  {w} -> {m}")
            vocab_data = {"version": "2.4", "words": words,
                          "mappings": mappings, "doc_hashes": [],
                          "strong": strong}

            for i, f in enumerate(files, 1):
                log(f"[{i}/{total}] 处理: {Path(f).name}")
                try:
                    processor = get_processor(f)
                    if not processor.load(f):
                        raise RuntimeError("文档加载失败")
                    count = processor.replace_all(mappings, mode="encrypt")

                    ext = Path(f).suffix
                    enc_doc = os.path.join(out_dir, f"{Path(f).stem}_encrypted{ext}")
                    if strategy == "skip" and os.path.exists(enc_doc):
                        log("  目标已存在，跳过")
                        skip_count += 1
                        continue
                    if not processor.save(enc_doc):
                        raise RuntimeError("加密文档保存失败")

                    vocab_data["doc_hashes"].append(compute_file_sha256(enc_doc))
                    ok_count += 1
                    log(f"  ✓ 替换 {count} 处 -> {Path(enc_doc).name}")
                except Exception as e:
                    fail_list.append(f"{Path(f).name}: {e}")
                    log(f"  ✗ 失败: {e}")

            if ok_count:
                enc_vocab = os.path.join(out_dir, "_batch_vocab.enc")
                with open(enc_vocab, "wb") as wf:
                    wf.write(CryptoEngine.encrypt_vocab(vocab_data, password))
                log(f"共享词表: {enc_vocab}（含 {len(vocab_data['doc_hashes'])} 个文档指纹）")

        log(f"批量加密完成：成功 {ok_count}，跳过 {skip_count}，失败 {len(fail_list)}")
        if fail_list:
            log("失败明细:")
            for item in fail_list:
                log(f"  - {item}")
        self.task.q.put(("result", (ok_count, skip_count, fail_list)))

    # ---------- 批量解密 ----------
    def batch_decrypt(self):
        src_dir = self.dec_src_var.get().strip()
        if not src_dir or not os.path.isdir(src_dir):
            messagebox.showwarning("提示", "请选择有效的解密源目录")
            return
        vocab_dir = self.dec_vocab_var.get().strip() or src_dir
        out_dir = self.dec_out_var.get().strip() or src_dir
        password = self.pwd_entry.get()
        if not password:
            messagebox.showwarning("提示", "请输入解密密码")
            return
        if self.task and self.task.is_running():
            messagebox.showwarning("提示", "正在处理中，请稍候...")
            return

        enc_files = [f for f in self.list_office_files(src_dir) if "_encrypted" in Path(f).name]
        if not enc_files:
            messagebox.showinfo("提示", "目录中没有找到 _encrypted 加密文档")
            return

        shared_vocab = os.path.join(vocab_dir, "_batch_vocab.enc")
        has_shared = os.path.exists(shared_vocab)

        # 预计算输出目标，检查覆盖
        targets = []
        for f in enc_files:
            targets.append(os.path.join(out_dir, restore_doc_name(Path(f).name)))
        strategy = check_overwrite_conflicts(targets, single=False)
        if strategy == "cancel":
            return

        self.batch_enc_btn.config(state=tk.DISABLED)
        self.batch_dec_btn.config(state=tk.DISABLED)
        self.log("=" * 40)
        self.log(f"批量解密开始：{len(enc_files)} 个文件" +
                 ("（共享词表模式）" if has_shared else "（独立词表模式）"))

        self.task = TaskThread(
            self.app.root,
            on_message=self.log,
            on_done=lambda err: self._batch_done(err, "decrypt", enc_files, out_dir)
        )
        self.task.start(self._batch_decrypt_worker, password, enc_files, vocab_dir, out_dir,
                        strategy, has_shared, shared_vocab)

    def _batch_decrypt_worker(self, password, enc_files, vocab_dir, out_dir,
                              strategy, has_shared, shared_vocab):
        log = lambda m: self.task.q.put(("log", m))
        os.makedirs(out_dir, exist_ok=True)
        total = len(enc_files)
        ok_count = 0
        skip_count = 0
        fail_list = []

        if has_shared:
            # ---- 共享词表模式 ----
            log("正在解密共享词表...")
            with open(shared_vocab, "rb") as f:
                vocab_data = CryptoEngine.decrypt_vocab(f.read(), password)
            words = vocab_data.get("words", [])
            mappings = vocab_data.get("mappings", {})
            doc_hashes = vocab_data.get("doc_hashes", [])
            log(f"词表解密成功，包含 {len(words)} 个词汇，{len(doc_hashes)} 个文档指纹")
            log("词表类型: " + ("强加密（多变体变长替换）" if vocab_data.get("strong") else "标准加密"))

            for i, f in enumerate(enc_files, 1):
                log(f"[{i}/{total}] 处理: {Path(f).name}")
                try:
                    doc_hash = compute_file_sha256(f)
                    if doc_hashes and doc_hash not in doc_hashes:
                        log("  ✗ 文档指纹不匹配，跳过")
                        fail_list.append(f"{Path(f).name}: 文档与共享词表不匹配")
                        continue

                    processor = get_processor(f)
                    if not processor.load(f):
                        raise RuntimeError("文档加载失败")
                    count = processor.replace_all(mappings, mode="decrypt")

                    dec_doc = os.path.join(out_dir, restore_doc_name(Path(f).name))
                    if strategy == "skip" and os.path.exists(dec_doc):
                        log("  目标已存在，跳过")
                        skip_count += 1
                        continue
                    if not processor.save(dec_doc):
                        raise RuntimeError("解密文档保存失败")
                    ok_count += 1
                    log(f"  ✓ 恢复 {count} 处 -> {Path(dec_doc).name}")
                except Exception as e:
                    fail_list.append(f"{Path(f).name}: {e}")
                    log(f"  ✗ 失败: {e}")
        else:
            # ---- 独立词表模式：按文件名配对 ----
            for i, f in enumerate(enc_files, 1):
                log(f"[{i}/{total}] 处理: {Path(f).name}")
                try:
                    orig_name = restore_doc_name(Path(f).name)
                    orig_stem = Path(orig_name).stem
                    vocab_path = os.path.join(vocab_dir, f"{orig_stem}_vocab.enc")
                    if not os.path.exists(vocab_path):
                        raise RuntimeError(f"未找到配套词表: {Path(vocab_path).name}")

                    with open(vocab_path, "rb") as vf:
                        vocab_data = CryptoEngine.decrypt_vocab(vf.read(), password)
                    mappings = vocab_data.get("mappings", {})
                    doc_hash = vocab_data.get("doc_hash")

                    if doc_hash:
                        actual = compute_file_sha256(f)
                        if doc_hash != actual:
                            raise RuntimeError("文档与词表不匹配")

                    processor = get_processor(f)
                    if not processor.load(f):
                        raise RuntimeError("文档加载失败")
                    count = processor.replace_all(mappings, mode="decrypt")

                    dec_doc = os.path.join(out_dir, orig_name)
                    if strategy == "skip" and os.path.exists(dec_doc):
                        log("  目标已存在，跳过")
                        skip_count += 1
                        continue
                    if not processor.save(dec_doc):
                        raise RuntimeError("解密文档保存失败")
                    ok_count += 1
                    log(f"  ✓ 恢复 {count} 处 -> {Path(dec_doc).name}")
                except Exception as e:
                    fail_list.append(f"{Path(f).name}: {e}")
                    log(f"  ✗ 失败: {e}")

        log(f"批量解密完成：成功 {ok_count}，跳过 {skip_count}，失败 {len(fail_list)}")
        if fail_list:
            log("失败明细:")
            for item in fail_list:
                log(f"  - {item}")
        self.task.q.put(("result", (ok_count, skip_count, fail_list)))

    def _batch_done(self, err, mode, files, out_dir):
        self.batch_enc_btn.config(state=tk.NORMAL)
        self.batch_dec_btn.config(state=tk.NORMAL)
        if err is None:
            self.log("全部完成。")
            # 读取 worker 末尾放入的批量统计
            result = None
            if hasattr(self.task, "results"):
                for kind, payload in self.task.results:
                    if kind == "result":
                        result = payload
            if result:
                ok_count, skip_count, fail_list = result
                action = "加密" if mode == "encrypt" else "解密"
                messagebox.showinfo("完成",
                    f"批量{action}完成！\n\n"
                    f"成功: {ok_count}\n跳过: {skip_count}\n失败: {len(fail_list)}\n"
                    f"输出目录: {out_dir}")
            else:
                messagebox.showinfo("完成", "批量处理完成！")
        else:
            exc, tb = err
            self.log(f"批量处理失败: {exc}")
            messagebox.showerror("错误", f"批量处理失败: {exc}")


class MainApp:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Office敏感词加解密工具 v2.4")
        self.root.geometry("1120x900")
        try:
            self.root.state("zoomed")   # Windows 最大化；其它平台不支持时忽略
        except tk.TclError:
            pass
        self.root.minsize(1080, 720)   # 固定最小尺寸，保证所有控件始终完整可见
        self._apply_window_icon()

        self.style = ttk.Style()
        self.style.theme_use("clam")

        # ========== Word 默认风格主题（Office 现代界面） ==========
        BG = "#F3F3F3"            # 窗口背景（浅灰白，同 Word 功能区外背景）
        PANEL = "#FFFFFF"         # 面板/内容区背景（白色）
        BTN = "#FFFFFF"           # 普通按钮默认（白底灰边，同 Word 工具栏）
        BTN_HOVER = "#D5E4F7"     # 按钮悬停（浅蓝，同 Office hover）
        BTN_PRESSED = "#BBD8F5"   # 按钮按下（中蓝）
        ACCENT = "#0078D4"        # Office 蓝（主按钮/焦点/选中）
        ACCENT_ACTIVE = "#1A86E0" # 主按钮悬停
        ACCENT_PRESSED = "#005A9E"# 主按钮按下
        BORDER = "#D1D1D1"        # 控件边框（浅灰）
        FG = "#1F1F1F"            # 主文字（Word 正文黑）
        FG_SUB = "#666666"        # 次级文字（灰）

        self.root.configure(bg=BG)
        self.style.configure(".", background=BG, foreground=FG)
        self.style.configure("TFrame", background=BG)
        self.style.configure("TLabel", background=BG, foreground=FG)
        self.style.configure("TLabelframe", background=BG, bordercolor="#D9D9D9",
                             relief="solid", borderwidth=1)
        self.style.configure("TLabelframe.Label", background=BG, foreground="#2B579A",
                             font=("Microsoft YaHei", 10, "bold"))
        self.style.configure("TNotebook", background=BG, bordercolor="#D9D9D9",
                             tabmargins=(8, 4, 8, 0))
        self.style.configure("TNotebook.Tab", background="#E8E8E8", foreground="#444444",
                             padding=(16, 6), font=("Microsoft YaHei", 10))
        self.style.map("TNotebook.Tab",
                       background=[("selected", "#FFFFFF"), ("active", "#D5E4F7")],
                       foreground=[("selected", "#1F1F1F"), ("active", "#2B579A")])
        # 普通按钮：白底灰边（Word 工具栏按钮），悬停浅蓝
        self.style.configure("TButton", background=BTN, foreground=FG,
                             bordercolor=BORDER, padding=(10, 4),
                             font=("Microsoft YaHei", 9))
        self.style.map("TButton",
                       background=[("active", BTN_HOVER),
                                   ("pressed", BTN_PRESSED),
                                   ("disabled", "#F2F2F2")],
                       foreground=[("active", "#2B579A"),
                                   ("pressed", "#1F3864"),
                                   ("disabled", "#A6A6A6")],
                       bordercolor=[("active", "#A7C9F0"),
                                    ("pressed", "#7BA9E4"),
                                    ("disabled", "#E8E8E8")])
        # 主操作按钮：Office 蓝（同 Word 对话框主按钮）
        self.style.configure("Accent.TButton", background=ACCENT, foreground="#FFFFFF",
                             bordercolor=ACCENT, padding=(10, 4),
                             font=("Microsoft YaHei", 9))
        self.style.map("Accent.TButton",
                       background=[("active", ACCENT_ACTIVE),
                                   ("pressed", ACCENT_PRESSED),
                                   ("disabled", "#A6C8E8")],
                       foreground=[("disabled", "#F0F0F0")])
        self.style.configure("TEntry", fieldbackground=PANEL, foreground=FG,
                             bordercolor=BORDER, insertcolor=FG, padding=3)
        self.style.map("TEntry", bordercolor=[("focus", ACCENT)])
        self.style.configure("TCheckbutton", background=BG, foreground=FG)
        self.style.map("TCheckbutton", background=[("active", BG)])
        self.style.configure("TScrollbar", background="#E1E1E1", troughcolor=BG,
                             bordercolor="#E1E1E1", arrowcolor="#555555")

        # 程序启动时自动释放说明书文件
        ensure_manual_files()

        # 打包环境：首次启动安装桌面快捷方式 + 右键菜单（路径变化时自动修正）
        ensure_shell_integration()

        self.build_ui()

        # 命令行/右键菜单传入的文件：自动填充对应面板并切到单文件页
        self.cli_mode, self.cli_path = parse_cli_args()
        self._apply_cli()

    def _apply_window_icon(self):
        """设置窗口/任务栏图标（优先打包内置 icon.ico，失败时静默忽略）"""
        try:
            if getattr(sys, "frozen", False):
                icon_path = os.path.join(sys._MEIPASS, "icon.ico")
            else:
                icon_path = os.path.join(get_app_dir(), "icon.ico")
            if os.path.exists(icon_path):
                self.root.iconbitmap(icon_path)
        except Exception:
            pass

    def _apply_cli(self):
        """右键菜单 / 命令行参数：自动选中文档、切换到单文件处理页"""
        if not self.cli_path:
            return
        try:
            if self.cli_mode == "decrypt":
                self.decrypt_frame.set_doc(self.cli_path)
                self.notebook.select(0)
                self.decrypt_frame.pwd_entry.focus_set()
            else:
                self.encrypt_frame.set_doc(self.cli_path)
                self.notebook.select(0)
                self.encrypt_frame.pwd_entry.focus_set()
        except Exception:
            pass

    def build_ui(self):
        # ========== 顶部标题区 ==========
        header = ttk.Frame(self.root, padding=6)
        header.pack(fill=tk.X)

        title_frame = ttk.Frame(header)
        title_frame.pack(fill=tk.X)

        title = ttk.Label(title_frame, text="Office文档敏感词加解密工具",
                         font=("Microsoft YaHei", 16, "bold"))
        title.pack(side=tk.LEFT)

        copyright_label = ttk.Label(
            title_frame,
            text="  版权所有 © Sam Li  |  邮箱: samstay@sina.com  |  v2.4",
            font=("Microsoft YaHei", 9),
            foreground="#666666"
        )
        copyright_label.pack(side=tk.LEFT, padx=(10, 0), pady=5)

        desc = ttk.Label(header,
                        text="支持Word(.docx)、Excel(.xlsx)、PowerPoint(.pptx) 的敏感词替换加密 · AES-256-GCM 词表保护 · 强加密多变体替换 · 支持批量处理",
                        font=("Microsoft YaHei", 9), foreground="#666666")
        desc.pack(pady=5, anchor=tk.W)

        # ========== 主内容区（固定布局，无滚动条） ==========
        main_frame = ttk.Frame(self.root, padding=6)
        main_frame.pack(fill=tk.BOTH, expand=True)

        main_frame.columnconfigure(0, minsize=300)
        main_frame.columnconfigure(1, weight=1)
        main_frame.rowconfigure(0, weight=1)

        # 左侧：词表管理（全局共享，两个 Tab 均使用）
        self.vocab_frame = VocabManagerFrame(main_frame, self)
        self.vocab_frame.grid(row=0, column=0, sticky="nsw", padx=(0, 8), pady=5)

        # 右侧：单文件处理 / 批量处理 两个 Tab
        self.notebook = ttk.Notebook(main_frame)
        self.notebook.grid(row=0, column=1, sticky="nsew", padx=(0, 0), pady=5)

        # ---- Tab 1: 单文件处理（加密 | 解密 左右分栏） ----
        tab_single = ttk.Frame(self.notebook, padding=6)
        self.notebook.add(tab_single, text="  单文件处理  ")
        tab_single.columnconfigure(0, weight=1, minsize=360)
        tab_single.columnconfigure(1, weight=1, minsize=360)
        tab_single.rowconfigure(0, weight=1)

        self.encrypt_frame = EncryptFrame(tab_single, self)
        self.encrypt_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 4), pady=2)

        self.decrypt_frame = DecryptFrame(tab_single, self)
        self.decrypt_frame.grid(row=0, column=1, sticky="nsew", padx=(4, 0), pady=2)

        # ---- Tab 2: 批量处理 ----
        tab_batch = ttk.Frame(self.notebook, padding=6)
        self.notebook.add(tab_batch, text="  批量处理  ")

        self.batch_frame = BatchFrame(tab_batch, self)
        self.batch_frame.pack(fill=tk.BOTH, expand=True)

        # ========== 底部状态栏 ==========
        footer = ttk.Frame(self.root, padding=5)
        footer.pack(fill=tk.X, side=tk.BOTTOM)

        btn_frame = ttk.Frame(footer)
        btn_frame.pack(side=tk.LEFT)

        ttk.Button(btn_frame, text="📖 使用说明书", command=lambda: open_text_file("README.txt"),
                  width=14).pack(side=tk.LEFT, padx=3)
        ttk.Button(btn_frame, text="© 版权声明", command=lambda: open_text_file("LICENSE.txt"),
                  width=12).pack(side=tk.LEFT, padx=3)
        ttk.Button(btn_frame, text="❤ 捐赠开发者",
                  command=lambda: webbrowser.open("https://samstay.dpdns.org/#donate"),
                  width=14).pack(side=tk.LEFT, padx=3)

        ttk.Label(footer, text="安全提示: 请妥善保管解密密码，密码丢失将无法恢复原文！",
                 foreground="red", font=("Microsoft YaHei", 9)).pack(side=tk.LEFT, padx=(20, 0))

        ttk.Button(footer, text="退出", command=self.root.quit, width=8).pack(side=tk.RIGHT, padx=5)

    def update_idletasks(self):
        self.root.update_idletasks()

    def run(self):
        self.root.mainloop()


if __name__ == "__main__":
    if "--gen-files" in sys.argv[1:]:
        # 打包脚本专用：只释放说明书/版权文件后退出，不启动界面
        ensure_manual_files()
        sys.exit(0)
    app = MainApp()
    app.run()
